"""Verifica los ejemplos didácticos en el Word general y su exportación real."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
import json
checks=[]
def test(name,fn):
    try:fn();checks.append({'name':name,'ok':True});print('PASS',name)
    except Exception as e:checks.append({'name':name,'ok':False,'error':str(e)});print('FAIL',name,str(e))
def truth(x):assert x
base=Document('Plantilla_Maestra_APA7_Word.docx')
export=Document('pruebas/apa-general-exportada.docx')
text='\n'.join(p.text for p in base.paragraphs)
test('APA general: ejemplos e instrucciones identificados, sin fuentes ficticias presentadas como reales',lambda:truth('ejemplos didácticos' in text and 'no una cita real' in text and 'no son fuentes reales' in text))
test('APA general: tabla y figura reales con título y nota',lambda:truth(len(base.tables)==1 and len(base.inline_shapes)==1 and 'Ejemplo de organización de actividades' in text and 'Secuencia de preparación de un trabajo' in text))
def runin():
    for level in [4,5]:
        p=next(p for p in base.paragraphs if p.style.name=='APA Nivel '+str(level))
        assert p.runs[0].bold is True and p.runs[-1].bold is False
        assert p.runs[0].italic==(level==5) and p.runs[-1].italic is False
        assert p.runs[0].text.endswith('.')
test('Niveles 4 y 5: texto en la misma línea sin negrita/cursiva de encabezado',runin)
def formats():
    for d in [base,export]:
        s=d.sections[0];assert [s.page_width.inches,s.page_height.inches,s.top_margin.inches,s.left_margin.inches]==[8.5,11,1,1]
        p=d.styles['APA Párrafo'];assert p.font.name=='Times New Roman' and p.font.size.pt==12 and p.paragraph_format.line_spacing==2
        assert len(d.tables)==1 and len(d.inline_shapes)==1
    with ZipFile('pruebas/apa-general-exportada.docx') as z:assert z.testzip() is None and 'word/document.xml' in z.namelist()
test('Word descargable y exportado: Carta, márgenes, tipografía, tabla e imagen',formats)
def table():
    for i,row in enumerate(base.tables[0].rows):
        for cell in row.cells:
            for side in ['left','right']:assert cell._tc.get_or_add_tcPr().find(qn('w:tcBorders')).find(qn('w:'+side)).get(qn('w:val'))=='nil'
            if i==0:assert all(r.bold for p in cell.paragraphs for r in p.runs)
    assert base.tables[0].rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
test('Tabla didáctica sin rejilla vertical y con encabezado repetible',table)
def models():
    files=list(Path('plantillas').glob('*.docx'));assert len(files)==14
    for file in files:
        d=Document(file);body='\n'.join(p.text for p in d.paragraphs)
        assert 'Esta APA general enseña el formato con ejemplos didácticos' not in body
        assert 'Secuencia de preparación de un trabajo' not in body
        assert 'word/media/image1.png' not in ZipFile(file).namelist()
test('Los 14 modelos conservan sus estructuras sin los ejemplos de la APA general',models)
def compact_table():
    assert base.styles['APA Número de elemento'].paragraph_format.keep_with_next
    assert base.styles['APA Título de elemento'].paragraph_format.keep_with_next
    for d in [base,export]:
        for i,row in enumerate(d.tables[0].rows):
            assert row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None
            for cell in row.cells:
                for p in cell.paragraphs:
                    assert bool(p.paragraph_format.keep_with_next)==(i<len(d.tables[0].rows)-1)
test('Tabla breve: título unido y filas conservadas juntas en Word general y exportado',compact_table)
def guide_media():
    with ZipFile('Guia_APA7_Estudio.docx') as z:
        assert not any(n.startswith('word/media/') for n in z.namelist())
test('Guía sin imágenes huérfanas ni archivos visuales duplicados',guide_media)
Path('pruebas/general-word-resultados.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
if not all(x['ok'] for x in checks):raise SystemExit(1)
