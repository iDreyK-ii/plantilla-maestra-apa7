from pathlib import Path
from zipfile import ZipFile
from docx import Document
import json
ROOT=Path(__file__).resolve().parent.parent
cat=json.loads((ROOT/'proyecto/catalogo.json').read_text());results=[]
for m in cat['models']:
 try:
  filename=ROOT/'plantillas'/('Plantilla_'+m['id'].replace('-','_')+'.docx');d=Document(filename);z=ZipFile(filename);assert z.testzip() is None
  s=d.sections[0];assert [s.page_width.inches,s.page_height.inches,s.top_margin.inches,s.bottom_margin.inches,s.left_margin.inches,s.right_margin.inches]==[8.5,11,1,1,1,1]
  st=d.styles['APA Párrafo'];assert st.font.name=='Times New Roman' and st.font.size.pt==12 and st.paragraph_format.line_spacing==2 and st.paragraph_format.first_line_indent.inches==.5
  st=d.styles['APA Referencia'];assert st.paragraph_format.first_line_indent.inches==-.5 and st.paragraph_format.left_indent.inches==.5
  assert all('APA Nivel '+str(i) in d.styles for i in range(1,6))
  assert any(b'PAGE' in z.read(n) for n in z.namelist() if n.startswith('word/header') and n.endswith('.xml'))
  texts=[p.text for p in d.paragraphs];text='\n'.join(texts)
  assert all(s['name'] in texts or (s['kind']=='keywords' and any(t.startswith(s['name']+':') for t in texts)) for s in m['sections'])
  for section in m['sections']:
   if section['kind']=='keywords':
    p=next(p for p in d.paragraphs if p.text.startswith(section['name']+':'));assert p.style.name=='APA Párrafo' and p.runs[0].italic
   else:
    p=next(p for p in d.paragraphs if p.text==section['name']);assert p.style.name=='APA Nivel '+str(section['level'])
  assert 'Ficha de uso de la plantilla — retirar antes de entregar' in text
  assert 'No es un formato oficial' in text
  for sid in m['sources']:assert next(s['url'] for s in cat['sources'] if s['id']==sid) in text
  assert not any(s in text for s in ['function export','onclick=','Cancelar','Abrir impresión','<html','<script'])
  assert b'altChunk' not in z.read('word/document.xml')
  assert not any('vbaProject' in n for n in z.namelist())
  if m['toc']:assert b'TOC' in z.read('word/document.xml')
  if m['id'] in ['upeu-perfil','upeu-doctoral']:assert len(d.tables)==2
  results.append(dict(name=m['label']+': archivo válido, formato, jerarquía, instrucciones y fuentes',ok=True));print('PASS',m['label'])
 except Exception as e:results.append(dict(name=m['label'],ok=False,error=str(e)));print('FAIL',m['label'],e)
try:
 d=Document(ROOT/'pruebas/modelo-trabajo-final.docx');text='\n'.join(p.text for p in d.paragraphs)
 assert 'ULTIMA_REVISION_2026' in text
 assert 'Formula un objetivo' not in text and 'Ficha de uso' not in text and 'Ejemplo orientativo' not in text
 results.append(dict(name='Trabajo exportado: última revisión y ningún texto del tutor',ok=True));print('PASS Trabajo exportado: última revisión y ningún texto del tutor')
except Exception as e:results.append(dict(name='Trabajo exportado',ok=False,error=str(e)))
(ROOT/'pruebas/modelos-word-resultados.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
if not all(x['ok'] for x in results):raise SystemExit(1)
