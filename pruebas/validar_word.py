from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree
import json
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
results=[]
def check(name,fn):
    try:fn();results.append({'name':name,'ok':True});print('PASS',name)
    except Exception as e:results.append({'name':name,'ok':False,'error':str(e)});print('FAIL',name,str(e))
def eq(a,b):assert a==b,(a,b)
def truth(x):assert x
for filename in ['Plantilla_Maestra_APA7_Word.docx','pruebas/trabajo-exportado.docx']:
    d=Document(filename);z=ZipFile(filename);label=Path(filename).name
    check(label+': OOXML real sin altChunk HTML',lambda:truth('[Content_Types].xml' in z.namelist() and b'altChunk' not in z.read('word/document.xml')))
    sec=d.sections[0]
    check(label+': Carta y cuatro márgenes de una pulgada',lambda:eq([sec.page_width.inches,sec.page_height.inches,sec.top_margin.inches,sec.bottom_margin.inches,sec.left_margin.inches,sec.right_margin.inches],[8.5,11,1,1,1,1]))
    p=d.styles['APA Párrafo'];r=d.styles['APA Referencia']
    check(label+': Times New Roman 12 y párrafo doble',lambda:eq([p.font.name,p.font.size.pt,p.paragraph_format.line_spacing,p.paragraph_format.first_line_indent.inches],['Times New Roman',12,2,.5]))
    check(label+': Referencias con sangría francesa y doble',lambda:eq([r.paragraph_format.left_indent.inches,r.paragraph_format.first_line_indent.inches,r.paragraph_format.line_spacing],[.5,-.5,2]))
    check(label+': Numeración automática PAGE en encabezado',lambda:truth(any(b'PAGE' in z.read(n) for n in z.namelist() if n.startswith('word/header') and n.endswith('.xml'))))
    check(label+': Estilos reutilizables de encabezados 1–5',lambda:truth(all('APA Nivel '+str(i) in d.styles for i in range(1,6))))
    check(label+': Portada seguida de salto real',lambda:truth(any(p.paragraph_format.page_break_before for p in d.paragraphs)))
    if filename.startswith('pruebas/'):
        check('Trabajo: tabla editable y celdas originales',lambda:eq([[c.text for c in row.cells] for row in d.tables[0].rows],[['Grupo','Participantes'],['A','24'],['B','28']]))
        check('Trabajo: imagen incorporada y texto alternativo',lambda:truth(len(d.inline_shapes)==1 and any(n.startswith('word/media/') for n in z.namelist()) and b'Diagrama de prueba' in z.read('word/document.xml')))
        check('Trabajo: cita y contenido real, no plantilla',lambda:truth(any('(García, 2026, p. 25)' in p.text for p in d.paragraphs) and not any('[Reemplaza' in p.text for p in d.paragraphs)))
        check('Trabajo: listas numeradas y con viñetas',lambda:truth('word/numbering.xml' in z.namelist() and b'numPr' in z.read('word/document.xml')))
        check('Trabajo: campo TOC actualizable',lambda:truth(b'TOC' in z.read('word/document.xml')))
        h4=next(p for p in d.paragraphs if p.text.startswith('Cuarto nivel.'))
        check('Trabajo: nivel 4 integrado conserva negrita solo en encabezado',lambda:eq([h4.runs[0].bold,h4.runs[-1].bold],[True,False]))
        h5=next(p for p in d.paragraphs if p.text.startswith('Quinto nivel.'))
        check('Trabajo: nivel 5 integrado conserva cursiva solo en encabezado',lambda:eq([h5.runs[0].italic,h5.runs[-1].italic],[True,False]))
Path('pruebas/word-estructura.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
if not all(r['ok'] for r in results):raise SystemExit(1)
