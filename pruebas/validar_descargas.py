from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from docx import Document
import json
ROOT=Path(__file__).resolve().parent.parent
checks=[]
z=ZipFile(ROOT/'Plantillas_APA7_Word.zip')
assert z.testzip() is None
files=[n for n in z.namelist() if n.endswith('.docx')]
assert len(files)==16 and 'Plantilla_Maestra_APA7_Word.docx' in files
for name in files:
 try:
  raw=z.read(name);assert raw==(ROOT/name).read_bytes()
  doc=Document(BytesIO(raw));sec=doc.sections[0]
  assert [sec.page_width.inches,sec.page_height.inches,sec.top_margin.inches,sec.bottom_margin.inches,sec.left_margin.inches,sec.right_margin.inches]==[8.5,11,1,1,1,1]
  p=doc.styles['APA Párrafo'];r=doc.styles['APA Referencia']
  assert [p.font.name,p.font.size.pt,p.paragraph_format.line_spacing,p.paragraph_format.first_line_indent.inches]==['Times New Roman',12,2,.5]
  assert [r.paragraph_format.left_indent.inches,r.paragraph_format.first_line_indent.inches,r.paragraph_format.line_spacing]==[.5,-.5,2]
  assert all('APA Nivel '+str(i) in doc.styles for i in range(1,6))
  xml=ZipFile(BytesIO(raw));assert any(b'PAGE' in xml.read(n) for n in xml.namelist() if n.startswith('word/header') and n.endswith('.xml'))
  assert not any('vbaProject' in n for n in xml.namelist())
  checks.append(dict(name=name+': formato e integridad dentro del paquete',ok=True))
 except Exception as e:checks.append(dict(name=name,ok=False,error=str(e)))
try:
 doc=Document(ROOT/'pruebas/trabajo-desde-descargas.docx');text='\n'.join(p.text for p in doc.paragraphs)
 assert 'ULTIMO_TEXTO_PROPIO_NO_PLANTILLA' in text
 assert 'Archivo listo para guardar' not in text and 'Guardar archivo preparado' not in text
 assert 'Tu trabajo actual' not in text and 'Revisa la carpeta Descargas' not in text
 checks.append(dict(name='Trabajo propio exportado sin avisos de descarga ni contenido sustituido',ok=True))
except Exception as e:checks.append(dict(name='Trabajo propio exportado',ok=False,error=str(e)))
(ROOT/'pruebas/descargas-word-resultados.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
for c in checks:print('PASS' if c['ok'] else 'FAIL',c['name'])
if not all(c['ok'] for c in checks):raise SystemExit(1)
