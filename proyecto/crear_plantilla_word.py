"""APA general didáctica .docx, estilos OOXML y ejemplos que se deben sustituir."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import base64
ROOT=Path(__file__).resolve().parent.parent

def flag(parent,tag,value=None):
    el=OxmlElement(tag)
    if value is not None: el.set(qn('w:val'),str(value))
    parent.append(el)
    return el

def style(doc,name,align=WD_ALIGN_PARAGRAPH.LEFT,indent=.5,hanging=False,bold=False,italic=False,level=None):
    s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.base_style=doc.styles['Normal'];s.next_paragraph_style=doc.styles['Normal']
    s.font.name='Times New Roman';s.font.size=Pt(12);s.font.bold=bold;s.font.italic=italic
    p=s.paragraph_format;p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=2
    p.first_line_indent=Inches(-.5 if hanging else indent);p.left_indent=Inches(.5 if hanging else 0)
    p.widow_control=True
    if level is not None:
        flag(s.element.get_or_add_pPr(),'w:outlineLvl',level-1);p.keep_with_next=True
    s.quick_style=True;s.priority=10 if level is None else 10+level
    return s

def build():
    d=Document();sec=d.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.5)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_before=normal.paragraph_format.space_after=Pt(0)
    normal.paragraph_format.first_line_indent=Inches(.5);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.widow_control=True
    style(d,'APA Párrafo')
    style(d,'APA Sin sangría',indent=0)
    style(d,'APA Título',WD_ALIGN_PARAGRAPH.CENTER,0,bold=True)
    style(d,'APA Subtítulo',WD_ALIGN_PARAGRAPH.CENTER,0)
    style(d,'APA Subtítulo 1',indent=0,bold=True)
    style(d,'APA Subtítulo 2',indent=0,bold=True,italic=True)
    for level in range(1,6):
        style(d,f'APA Nivel {level}',WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT,.5 if level>=4 else 0,bold=True,italic=level in (3,5),level=level)
    style(d,'APA Referencia',hanging=True)
    quote=style(d,'APA Cita larga',indent=0);quote.paragraph_format.left_indent=Inches(.5)
    style(d,'APA Número de elemento',indent=0,bold=True).paragraph_format.keep_with_next=True
    style(d,'APA Título de elemento',indent=0,italic=True).paragraph_format.keep_with_next=True
    style(d,'APA Nota de elemento',indent=0)
    style(d,'APA Portada',WD_ALIGN_PARAGRAPH.CENTER,0)
    # Numeración de campo PAGE: Word recalcula al cambiar el documento.
    header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent=Inches(0);header.paragraph_format.line_spacing=1
    run=header.add_run();run.font.name='Times New Roman';run.font.size=Pt(12)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r);run._r.addnext(field)
    sec.different_first_page_header_footer=False
    flag(d.settings.element,'w:updateFields','true')
    # Portada. Espaciado real, no párrafos en blanco para simular posición.
    p=d.add_paragraph('[Título del trabajo académico]','APA Título')
    p.paragraph_format.space_before=Pt(96);p.paragraph_format.space_after=Pt(24)
    for text in ['[Nombre del estudiante]','[Institución / Universidad]','[Facultad o escuela, si tu institución la solicita]','[Código y nombre de la asignatura]','[Nombre del docente]','[Fecha de entrega]']:
        d.add_paragraph(text,'APA Portada')
    p=d.add_paragraph('[Título del trabajo académico]','APA Título');p.paragraph_format.page_break_before=True
    from ejemplos_build import append_word_examples
    append_word_examples(d,ROOT)
    d.core_properties.title='PLANTILLA MAESTRA APA 7 — Trabajo de estudiante'
    d.core_properties.subject='Plantilla Word editable, Carta, márgenes 2,54 cm y estilos APA'
    d.core_properties.author='Plantilla Maestra APA 7'
    d.core_properties.comments='Base APA 7 para estudiantes. Adaptar estructura y requisitos institucionales; eliminar los textos entre corchetes antes de entregar.'
    out=ROOT/'Plantilla_Maestra_APA7_Word.docx';d.save(out)
    payload=base64.b64encode(out.read_bytes()).decode()
    (ROOT/'proyecto/plantilla-word.js').write_text('// Plantilla DOCX incorporada para descargar sin conexión. Generada por crear_plantilla_word.py.\nconst WORD_TEMPLATE_BASE64='+repr(payload)+';\nfunction downloadWordTemplate(){const data=Uint8Array.from(atob(WORD_TEMPLATE_BASE64),c=>c.charCodeAt(0));if(!download("Plantilla_Maestra_APA7_Word.docx",data,"application/vnd.openxmlformats-officedocument.wordprocessingml.document"))return;toast("Plantilla Word lista. Revisa Descargas o usa el enlace para guardarla; reemplaza los corchetes.")}\n')
    print(out, out.stat().st_size, 'bytes')
if __name__=='__main__':build()
