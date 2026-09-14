"""Precompila cada modelo: la descarga de la plantilla no necesita generarse ni conectarse a internet."""
import json,base64,unicodedata
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent.parent
cat=json.loads((ROOT/'proyecto/catalogo.json').read_text())
outdir=ROOT/'plantillas';outdir.mkdir(exist_ok=True)
base=ROOT/'Plantilla_Maestra_APA7_Word.docx'
files={}
def para(d,text,style='APA Párrafo',pagebreak=False):
 p=d.add_paragraph(text,style);p.paragraph_format.page_break_before=pagebreak;return p
def table(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,h in enumerate(headers):
  t.rows[0].cells[i].text=h
  for p in t.rows[0].cells[i].paragraphs:
   p.style=d.styles['APA Sin sangría']
   for r in p.runs:r.bold=True
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):
   cells[i].text=x
   for p in cells[i].paragraphs:p.style=d.styles['APA Sin sangría']
 for ri,row in enumerate(t.rows):
  for cell in row.cells:
   borders=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if (side=='top' and ri==0) or (side=='bottom' and ri in [0,len(t.rows)-1]) else 'nil');e.set(qn('w:sz'),'6');borders.append(e)
   cell._tc.get_or_add_tcPr().append(borders)
 return t
for m in cat['models']:
 d=Document(base)
 for child in list(d.element.body):
  if child.tag!=qn('w:sectPr'):d.element.body.remove(child)
 # Retirar la figura didáctica de la base: cada modelo conserva solo sus propios recursos.
 for rid,rel in list(d.part.rels.items()):
  if rel.reltype==RT.IMAGE:d.part.drop_rel(rid)
 if m['cover']:
  p=para(d,'[Título del trabajo]','APA Título');p.paragraph_format.space_before=Pt(96);p.paragraph_format.space_after=Pt(24)
  for txt in ['[Nombre del estudiante o autores]','Universidad Peruana Unión' if m['category']=='UPeU' else '[Institución]','[Facultad / Escuela / Programa]','[Asignatura o grado al que se opta]','[Docente o asesor, según corresponda]','[Fecha de entrega]']:para(d,txt,'APA Portada')
 else:
  para(d,'[Título del manuscrito]','APA Título')
  para(d,'[Autoría y afiliación solo si el destino las admite; para revisión anónima, eliminarlas]','APA Sin sangría')
 if m['toc']:
  para(d,'Índice','APA Título',m['cover']);p=para(d,'','APA Sin sangría');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-5" \\h \\z \\u');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='[En Word: clic derecho → Actualizar campo / tabla completa.]';r.append(t);field.append(r);p._p.append(field)
 for i,s in enumerate(m['sections']):
  if s['kind']=='keywords':
   p=para(d,'','APA Párrafo');r=p.add_run(s['name']+': ');r.italic=True;p.add_run('['+s['guide']+']');continue
  p=para(d,s['name'],'APA Nivel '+str(s['level']),s['key']=='referencias' or (i==0 and m['cover']))
  text='['+('Si corresponde a tu diseño o encargo: ' if s['optional'] else '')+s['guide']+']'
  para(d,text,'APA Referencia' if s['key']=='referencias' else 'APA Sin sangría' if s['kind']=='noindent' else 'APA Párrafo')
  if s['kind']=='schedule':table(d,['Actividad','Inicio','Fin'],[['[Actividad]','[Fecha]','[Fecha]'] for _ in range(3)])
  if s['kind']=='budget':table(d,['Recurso','Cantidad','Costo unitario','Total'],[['[Recurso]','[Cantidad]','[Importe]','[Importe]'] for _ in range(3)])
 para(d,'Ficha de uso de la plantilla — retirar antes de entregar','APA Nivel 1',True)
 para(d,m['label']+' · '+m['org'],'APA Sin sangría')
 para(d,'Adaptación educativa en formato APA 7. No es un formato oficial ni implica aprobación de la universidad. Los textos entre corchetes son indicaciones para reemplazar. Los apartados opcionales deben decidirse con tu asesor.','APA Sin sangría')
 for warning in m['warnings']:para(d,warning,'APA Sin sangría')
 para(d,'Fuentes de la estructura (no son referencias de tu investigación):','APA Sin sangría')
 for sid in m['sources']:
  src=next(x for x in cat['sources'] if x['id']==sid);para(d,src['org']+'. '+src['title']+'. '+src['date']+'. Consulta: '+src['checked']+'. '+src['url'],'APA Sin sangría')
 if not m['sources']:para(d,'Modelo didáctico propio, no atribuido a una universidad.','APA Sin sangría')
 d.core_properties.title='Plantilla guiada — '+m['label'];d.core_properties.author='Plantilla Maestra APA 7';d.core_properties.subject='Adaptación estructural con fuentes; no es documento oficial de la institución';d.core_properties.comments='Modelo '+m['id']+'; revisión '+cat['revision']+'; '+m['adaptation'];d.core_properties.keywords=m['id']
 filename='Plantilla_'+m['id'].replace('-','_')+'.docx';path=outdir/filename;d.save(path)
 files[m['id']]={'name':filename,'base64':base64.b64encode(path.read_bytes()).decode()}
(ROOT/'proyecto/modelos-word.js').write_text('// Modelos DOCX precompilados: descarga local inmediata.\nconst MODEL_WORD_FILES='+json.dumps(files,separators=(',',':'))+';\n')
print(len(files),'plantillas Word listas para descargar')
