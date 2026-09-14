from pathlib import Path
import json,base64
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent.parent
lessons=json.loads((ROOT/'proyecto/lecciones.json').read_text())
lessons[6][2]='Paráfrasis: (Apellido, año). Cita textual breve: «texto exacto» (Apellido, año, p. X). Sustituye los marcadores con una fuente real.'
lessons[7][2]='Estructura de un libro: Apellido, A. A. (Año). Título del libro en cursiva. Editorial. No es una referencia real para copiar.'
lessons[7][3]='1. Abre Biblioteca de referencias. 2. Elige el tipo de fuente. 3. Completa sus datos reales. 4. Inserta la cita y revisa la entrada generada.'
lessons[9][5]='Comprueba la legibilidad de la imagen en el Word exportado y añade texto alternativo. Revisa atribución, licencia y privacidad.'
lessons[12]=['Guardar, descargar y volver','Son tres acciones distintas: el guardado local conserva tu sesión; Word sirve para editar o entregar un documento; Copia APA conserva la biblioteca y los datos de esta aplicación.','Plantilla normal → archivo nuevo. Mi trabajo Word → tu contenido actual. Copia APA → respaldo para continuar aquí.','1. Pulsa Plantillas Word para un modelo vacío. 2. Usa Mi trabajo Word para exportar lo escrito. 3. Conserva además una Copia APA. Puedes abrir DOCX como copia nueva con revisión de formato.','Descargar el HTML de la aplicación pensando que es el documento Word, o borrar datos del navegador sin una copia.','Las plantillas del HTML incluyen enlaces de archivo directos. Si el navegador no inicia la descarga, usa Guardar en mi equipo cuando esté disponible, o el enlace alternativo del aviso. El DOCX se importa con límites; para recuperar fielmente los datos de la web, usa Copia APA.']
# Pasos breves y concretos para los módulos restantes.
steps={0:'1. Elige el tipo de trabajo. 2. Consulta la consigna de tu curso. 3. Separa tus ideas de las que provienen de fuentes. 4. Usa los estilos y registra cada fuente.',1:'1. Empieza con una plantilla APA. 2. Escribe sin agregar espacios para alinear. 3. Al abrir el Word, confirma Carta y márgenes de 2,54 cm.',2:'1. Abre Configuración. 2. Completa título y datos del curso. 3. Deja vacíos los campos que no correspondan. 4. Confirma si tu programa exige una portada propia.',3:'1. Coloca el cursor en el párrafo. 2. Aplica Párrafo APA. 3. Escribe una idea central y desarrolla su explicación. 4. Usa Enter para crear el siguiente párrafo.',4:'1. Usa Párrafo APA para el cuerpo. 2. Usa Referencia para cada entrada bibliográfica. 3. No añadas espacios ni tabulaciones al inicio.',5:'1. Decide si el apartado es principal o depende de otro. 2. Aplica el nivel correspondiente. 3. Revisa el esquema. 4. En los niveles 4–5, continúa el texto después del punto.',6:'1. Registra la fuente. 2. Elige una cita narrativa o parentética. 3. Añade el localizador si copias palabras exactas. 4. Comprueba la referencia.',8:'1. Inserta Tabla APA. 2. Completa número y título. 3. Introduce los datos en celdas editables. 4. Añade una nota solo si hace falta.',9:'1. Inserta Figura APA. 2. Carga una imagen local. 3. Completa número, título y texto alternativo. 4. Explica símbolos y atribuye la fuente cuando corresponda.',10:'1. Aplica encabezados al documento. 2. Comprueba el orden en el panel Documento. 3. Inserta el índice si tu encargo lo pide. 4. Actualiza el campo en Word antes de entregar.',11:'1. Ejecuta Revisar APA. 2. Lee cada aviso y su explicación. 3. Revisa argumentos, citas y ortografía por tu cuenta. 4. Abre el archivo descargado y comprueba el resultado.'}
for i,text in steps.items():lessons[i][3]=text
(ROOT/'proyecto/lecciones.json').write_text(json.dumps(lessons,ensure_ascii=False,indent=2))
quiz=[
 ('¿APA impone la misma estructura a todos los trabajos?',['Sí, siempre cinco capítulos','No: el género y la institución también importan'],1,'APA orienta formato y atribución. El esquema depende del tipo de trabajo y del encargo.'),
 ('¿Qué configuración tiene nuestra plantilla?',['Carta, márgenes 2,54 cm y doble espacio','A4 con márgenes de 1 cm'],0,'La base de esta aplicación usa Carta, Times New Roman 12 y doble espacio; APA admite otras fuentes.'),
 ('¿Cómo se numera la portada?',['Con el campo automático de página','Escribiendo un 1 a mano'],0,'El campo PAGE recalcula la numeración al cambiar el documento.'),
 ('¿Cómo empieza un párrafo de cuerpo?',['Con espacios escritos a mano','Con sangría real de primera línea'],1,'El estilo Párrafo APA aplica 1,27 cm de primera línea.'),
 ('¿Qué sangría corresponde a referencias?',['Francesa','Primera línea'],0,'La primera línea queda al margen; las siguientes entran 1,27 cm.'),
 ('¿Un nivel 3 depende normalmente de…?',['Un nivel 2','Cualquier estilo que se vea bonito'],0,'Los niveles expresan relaciones entre apartados, no decoración.'),
 ('¿Qué añades a una cita textual, además de autor y año?',['Un localizador: página u otra ubicación','Solo el enlace de la web'],0,'Una cita textual necesita poder localizarse en la fuente.'),
 ('¿Una referencia ficticia sirve para completar la lista?',['No: registra fuentes reales que utilizaste','Sí, si tiene el formato correcto'],0,'La presentación no reemplaza la veracidad ni la consulta de fuentes.'),
 ('¿Cómo entregas una tabla editable?',['Como una captura de pantalla','Como celdas con datos'],1,'Las celdas conservan la edición, búsqueda y accesibilidad del contenido.'),
 ('¿Qué ayuda a comprender una figura sin verla?',['El texto alternativo','El color del borde'],0,'Describe su información relevante, sin repetir innecesariamente todo el título.'),
 ('¿Todos los trabajos APA deben tener índice?',['Sí','No: depende del encargo'],1,'El índice es útil en documentos extensos, pero no es una exigencia APA universal.'),
 ('¿Cero avisos automáticos equivale a un trabajo perfecto?',['No; falta la revisión humana','Sí, ya está certificado'],0,'El revisor no verifica rigor, originalidad, exactitud de fuentes ni aprobación institucional.'),
 ('¿Qué copia conserva mejor la biblioteca de esta web?',['Una captura de pantalla','Copia APA (.json)'],1,'Word conserva contenido editable; Copia APA también guarda la biblioteca y la configuración de la aplicación.')]
q=[dict(q=a,options=b,correct=c,feedback=d) for a,b,c,d in quiz]
(ROOT/'proyecto/aprendizaje-data.js').write_text('const LEARNING_LESSONS='+json.dumps(lessons,ensure_ascii=False)+';\nconst STUDIO_QUIZ='+json.dumps(q,ensure_ascii=False)+';\n')
d=Document(ROOT/'Plantilla_Maestra_APA7_Word.docx')
for n in list(d.element.body):
 if n.tag!=qn('w:sectPr'):d.element.body.remove(n)
# No conservar imágenes huérfanas de la plantilla base en la guía sin figuras.
for rid,rel in list(d.part.rels.items()):
 if rel.reltype.endswith('/image'):d.part.drop_rel(rid)
d.add_paragraph('Aprende APA 7, paso a paso','APA Título');d.add_paragraph('Guía de estudio · Plantilla Maestra APA 7','APA Sin sangría');d.add_paragraph('Material educativo, no una certificación ni un reemplazo del manual APA o de las instrucciones institucionales. Los ejemplos con marcadores no son fuentes reales.','APA Párrafo')
for i,l in enumerate(lessons):
 p=d.add_paragraph(f'{i+1:02}. {l[0]}','APA Nivel 1');p.paragraph_format.page_break_before=True
 d.add_paragraph(l[1],'APA Párrafo')
 for label,text in [('Ejemplo orientativo',l[2]),('Hazlo paso a paso',l[3]),('Evita este error',l[4]),('Consejo',l[5]),('Comprueba lo aprendido',q[i]['q']),('Respuesta explicada',q[i]['feedback'])]:
  d.add_paragraph(label,'APA Nivel 2');d.add_paragraph(text,'APA Párrafo')
d.add_paragraph('Dónde consultar más','APA Nivel 1');d.add_paragraph('APA Style: https://apastyle.apa.org/style-grammar-guidelines/paper-format\nFuentes institucionales y modelos: consulta FUENTES_Y_MODELOS_APA7.md y las fichas dentro de la aplicación.\nNo se ha certificado la vigencia de todos los formatos institucionales para cada programa.','APA Sin sangría')
d.core_properties.title='Guía de estudio APA 7';d.core_properties.author='Plantilla Maestra APA 7';d.save(ROOT/'Guia_APA7_Estudio.docx')
(ROOT/'proyecto/guia-word.js').write_text('const STUDY_GUIDE_BASE64='+json.dumps(base64.b64encode((ROOT/'Guia_APA7_Estudio.docx').read_bytes()).decode())+';\n')
print('13 módulos revisados, 13 prácticas y guía Word listos')
