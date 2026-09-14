"""Misma guía didáctica para la APA general en Word y en el editor web."""
from pathlib import Path
import base64
import json


def general_data(root):
    data=json.loads((root/'proyecto/ejemplos_apa.json').read_text())
    data['figureData']='data:image/png;base64,'+base64.b64encode((root/'assets/ejemplo-esquema.png').read_bytes()).decode()
    return data


def append_word_examples(doc,root):
    from docx.shared import Inches
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    data=general_data(root)
    doc.add_paragraph('['+data['notice']+']','APA Sin sangría')
    for b in data['blocks']:
        kind=b['kind']
        if kind=='heading':
            p=doc.add_paragraph(b['text'],'APA Nivel '+str(b['level']));p.paragraph_format.page_break_before=bool(b.get('pagebreak'))
        elif kind=='runin':
            p=doc.add_paragraph('', 'APA Nivel '+str(b['level']));r=p.add_run(b['text']);r.bold=True;r.italic=b['level']==5
            r=p.add_run(b['body']);r.bold=False;r.italic=False
        elif kind=='instruction':doc.add_paragraph('[Guía: '+b['text']+']','APA Sin sangría')
        elif kind=='paragraph':doc.add_paragraph(b['text'],'APA Párrafo')
        elif kind=='quote':doc.add_paragraph(b['text'],'APA Cita larga')
        elif kind=='reference':
            p=doc.add_paragraph('', 'APA Referencia');p.add_run(b['before']);p.add_run(b['italic']).italic=True;p.add_run(b['after'])
        elif kind in ('table','figure'):
            doc.add_paragraph(b['number'],'APA Número de elemento');doc.add_paragraph(b['title'],'APA Título de elemento')
            if kind=='table':
                t=doc.add_table(rows=1,cols=len(b['headers']))
                for i,text in enumerate(b['headers']):t.rows[0].cells[i].text=text
                for row in b['rows']:
                    cells=t.add_row().cells
                    for i,text in enumerate(row):cells[i].text=text
                for ri,row in enumerate(t.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.style=doc.styles['APA Sin sangría']
                            p.paragraph_format.keep_with_next=ri<len(t.rows)-1
                            for r in p.runs:r.bold=ri==0
                        borders=OxmlElement('w:tcBorders')
                        for side in ['top','bottom','left','right']:
                            e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if (side=='top' and ri==0) or (side=='bottom' and ri in (0,len(t.rows)-1)) else 'nil');e.set(qn('w:sz'),'6');borders.append(e)
                        cell._tc.get_or_add_tcPr().append(borders)
                repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            else:
                shape=doc.add_picture(str(root/'assets'/b['file']),width=Inches(6));shape._inline.docPr.set('descr',b['alt'])
                doc.paragraphs[-1].style=doc.styles['APA Sin sangría']
            doc.add_paragraph(b['note'],'APA Nota de elemento')
